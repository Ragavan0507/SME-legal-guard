import streamlit as st
import pdfplumber
import docx
from openai import OpenAI
import json
import datetime
from fpdf import FPDF

# --- 1. IMPORT TEMPLATES ---
try:
    from templates import TEMPLATES
except ImportError:
    st.error("Critical Error: 'templates.py' not found! Please ensure it is in the same folder.")
    TEMPLATES = {"Error": "Template file missing."}

# --- 2. CONFIGURATION & API SETUP ---
st.set_page_config(page_title="SME Legal Guard", layout="wide", page_icon="⚖️")
# --- HIDE STREAMLIT BRANDING ---
hide_st_style = """
            <style>
            MainMenu {visibility: hidden;}
            footer {visibility: hidden;}
            header {visibility: hidden;}
            viewer-badge {display: none !important;}
            div[data-testid="stToolbar"] {display: none !important;}
            div[data-testid="stDecoration"] {display: none !important;}
            GithubIcon {visibility: hidden;}
            </style>
            """
st.markdown(hide_st_style, unsafe_allow_html=True)
# CLOUD DEPLOYMENT TIP: 
# For local testing, replace st.secrets["GITHUB_TOKEN"] with your "Actual_Token_String"
# For live deployment, use st.secrets for security.
try:
    GITHUB_TOKEN = st.secrets["GITHUB_TOKEN"]
except:
    GITHUB_TOKEN = "YOUR_GITHUB_TOKEN_HERE" # Fallback for local testing

client = OpenAI(
    base_url="https://models.inference.ai.azure.com",
    api_key=GITHUB_TOKEN,
)

# --- 3. HELPER FUNCTIONS ---

def extract_text(uploaded_file):
    """Extracts text from PDF, DOCX, and TXT files."""
    try:
        filename = uploaded_file.name.lower()
        if filename.endswith('.pdf'):
            with pdfplumber.open(uploaded_file) as pdf:
                return "".join([page.extract_text() for page in pdf.pages if page.extract_text()])
        elif filename.endswith('.docx'):
            doc = docx.Document(uploaded_file)
            return "\n".join([para.text for para in doc.paragraphs])
        elif filename.endswith('.txt'):
            return uploaded_file.read().decode("utf-8", errors="ignore")
        return None
    except Exception as e:
        st.error(f"Error reading file: {e}")
        return None

def clean_text_for_pdf(text):
    """Cleans text to prevent FPDF 'Latin-1' encoding errors."""
    if not text: return ""
    replacements = {
        '\u2013': '-', '\u2014': '-', '\u2018': "'", '\u2019': "'", 
        '\u201c': '"', '\u201d': '"', '\u2022': '*', '\u2026': '...'
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode('latin-1', 'ignore').decode('latin-1')

def create_pdf(data, timestamp, doc_name):
    """Generates a professional PDF audit report."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="SME Legal Guard: Risk & Negotiation Report", ln=True, align='C')
    pdf.set_font("Arial", size=10)
    pdf.cell(200, 10, txt=f"Document: {doc_name} | Date: {timestamp}", ln=True, align='C')
    pdf.ln(10)
    
    # Executive Summary
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(200, 10, txt="Executive Summary:", ln=True)
    pdf.set_font("Arial", size=11)
    pdf.multi_cell(0, 10, txt=clean_text_for_pdf(data.get('summary', '')))
    pdf.ln(5)
    
    # Risk Score
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(200, 10, txt=f"Contract Risk Score: {data.get('risk_score')}/100", ln=True)
    pdf.ln(5)
    
    # Detailed Risks
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(200, 10, txt="Clause Analysis & Negotiation Advice:", ln=True)
    for risk in data.get('risks', []):
        pdf.set_font("Arial", 'B', 11)
        pdf.cell(0, 10, txt=clean_text_for_pdf(f"- {risk.get('clause')} [{risk.get('level')}]"), ln=True)
        pdf.set_font("Arial", size=10)
        pdf.multi_cell(0, 7, txt=clean_text_for_pdf(f"Impact: {risk.get('impact')}"))
        pdf.set_text_color(0, 51, 102) # Blue color for tips
        pdf.multi_cell(0, 7, txt=clean_text_for_pdf(f"Negotiation Tip: {risk.get('alternative')}"))
        pdf.set_text_color(0, 0, 0) # Reset to black
        pdf.ln(2)
        
    return pdf.output(dest='S').encode('latin-1')

# --- 4. SIDEBAR ---
if 'audit_log' not in st.session_state:
    st.session_state.audit_log = []

with st.sidebar:
    st.title("🛡️ SME Command Center")
    
    # Template Downloader
    st.subheader("📁 Safe Templates")
    st.caption("Standard SME-friendly drafts")
    template_choice = st.selectbox("Select Template", list(TEMPLATES.keys()))
    st.download_button(
        label=f"📥 Download {template_choice}",
        data=TEMPLATES[template_choice],
        file_name=f"{template_choice.replace(' ', '_')}.txt",
        mime="text/plain"
    )
    
    st.markdown("---")
    
    # Audit Log
    st.subheader("📋 Session Audit Log")
    if not st.session_state.audit_log:
        st.write("No active scans.")
    else:
        for entry in st.session_state.audit_log:
            st.write(f"✅ {entry}")

# --- 5. MAIN UI ---
st.title("⚖️ SME Legal Guard")
st.markdown("### AI Risk Auditor, Negotiator & Template Hub")
st.write("Upload your contract to identify hidden risks and get professional renegotiation tips.")

uploaded_file = st.file_uploader("Upload Contract (PDF, DOCX, TXT)", type=["pdf", "docx", "txt"])

if uploaded_file:
    text_content = extract_text(uploaded_file)
    
    if text_content:
        st.success(f"Successfully loaded: {uploaded_file.name}")

        if st.button("🚀 Start Deep AI Audit"):
            with st.spinner("AI Lawyer is reviewing your document..."):
                current_time = datetime.datetime.now()
                timestamp_str = current_time.strftime("%Y-%m-%d %H:%M:%S")
                
                # Master Prompt
                prompt = f"""
                Analyze the following Indian contract text. Return ONLY a valid JSON object:
                {{
                    "risk_score": (integer 1-100),
                    "summary": "2-sentence summary of the document",
                    "risks": [
                        {{
                            "clause": "Name", 
                            "level": "High/Medium/Low", 
                            "impact": "English explanation", 
                            "hindi": "Simple Hindi explanation",
                            "alternative": "Suggested renegotiation wording"
                        }}
                    ]
                }}
                Text: {text_content[:7000]}
                """

                try:
                    response = client.chat.completions.create(
                        messages=[{"role": "user", "content": prompt}],
                        model="gpt-4o-mini"
                    )
                    
                    # Clean response for JSON
                    content = response.choices[0].message.content
                    if "```json" in content:
                        content = content.split("```json")[1].split("```")[0]
                    
                    data = json.loads(content.strip())
                    
                    # Update Audit Log
                    st.session_state.audit_log.append(f"{uploaded_file.name} - {current_time.strftime('%H:%M:%S')}")

                    # --- DISPLAY DASHBOARD ---
                    st.markdown("---")
                    c1, c2 = st.columns([1, 3])
                    
                    with c1:
                        st.metric("Risk Score", f"{data['risk_score']}/100")
                        if data['risk_score'] > 60: st.error("🛑 High Risk")
                        elif data['risk_score'] > 30: st.warning("⚠️ Medium Risk")
                        else: st.success("✅ Low Risk")

                    with c2:
                        st.subheader("Executive Summary")
                        st.write(data['summary'])

                    st.subheader("🚩 Clause Analysis & Negotiation Tips")
                    for risk in data['risks']:
                        with st.expander(f"{risk['clause']} — {risk['level']} Risk"):
                            st.write(f"**English:** {risk['impact']}")
                            st.write(f"**हिंदी:** {risk['hindi']}")
                            st.info(f"💡 **Negotiation Alternative:** {risk['alternative']}")

                    # --- PDF DOWNLOAD ---
                    pdf_bytes = create_pdf(data, timestamp_str, uploaded_file.name)
                    st.download_button(
                        label="📥 Download Full Audit Report (PDF)",
                        data=pdf_bytes,
                        file_name=f"SME_Audit_{uploaded_file.name}.pdf",
                        mime="application/pdf"
                    )
                    st.balloons()

                except Exception as e:
                    st.error(f"Analysis Error: {e}")
    else:
        st.warning("File is empty or unreadable.")

# --- FOOTER ---
st.markdown("---")

st.caption("Disclaimer: This tool is AI-powered and intended for guidance. It is not a substitute for legal advice.")

